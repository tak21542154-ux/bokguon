# -*- coding: utf-8 -*-
"""스캔 58라운드 회귀 — 개산 엔진(HIGH) + 수량식(MEDIUM) + 표준단가 단위불일치
   (MEDIUM) + 죽은 게이트 삭제(LOW) + 엑셀 정제 일관성(LOW).

  조사 에이전트가 실측(격리 DB + 대조군)으로 증명한 결함:

  ① _validate() 의 choice 검증이 화면이 실제로 보내는 값(문자열)과 항상 어긋났다
     — choices 가 숫자인 항목(가드레일 지주간격=[2,4])은 "2"/"4" 둘 다 400 으로
     거부되고, 빈칸으로 둬야만 통과해 기본값(4m)이 조용히 박혔다(실제 2m 현장이면
     26.5% 과소산출). list_patterns() 도 choices/type/min/max 를 화면에 안 내려줘서
     admin.js 가 애초에 선택지를 그릴 수 없었다.
  ② _posts() 의 지주 수 공식이 floor 였다(자기 docstring 은 ceil 이 맞다고 말한다) —
     L 이 지주간격의 배수가 아니면 항상 1개씩 부족했다.
  ③ _load_price_book() 이 DB 단가행의 unit 을 전혀 확인하지 않았다 — sku_key 는
     그대로 두고 unit 만 바꾸면(예: '개'→'m') 화면엔 옛 단위가 찍히는데 실제
     곱셈은 새 단가로 일어나 금액이 조용히 틀어진다. is_demo 필터·결정적 순서도 없었다.
  ④ _gate() 의 SPREAD_MAX 불확실성 게이트가 구조적으로 죽어 있었다(high/low 비율이
     항상 1.2 로 고정, 1.6 을 절대 못 넘음) — 삭제.
  ⑤ api_ops_kpi_export 만 형제 엑셀 내보내기들과 달리 _xl_safe 정제를 안 거쳤다.

      python test_scan_round58_fixes.py
"""
import io
import math
import os
import sys
import tempfile

HERE = os.path.dirname(os.path.abspath(__file__))
PASS, FAIL = [], []


def chk(cond, name, detail=""):
    (PASS if cond else FAIL).append((name, detail))
    print(("  [OK]   " if cond else "  [FAIL] ") + name + ("" if cond else f"   {detail}"))


def read(rel):
    return io.open(os.path.join(HERE, rel), encoding="utf-8").read()


APP_SRC = read("app.py")
RE_SRC = read("repair_estimate.py")
JS_SRC = read(os.path.join("static", "admin.js"))

print("═" * 72)
print("  스캔 58라운드 회귀 — 개산엔진 choice검증·수량식·단가단위·죽은게이트·엑셀정제")
print("═" * 72)


def _body(src, fn_name):
    start = src.index(f"\ndef {fn_name}(")
    end = src.find("\ndef ", start + 1)
    return src[start:end if end != -1 else None]


# ══════════════════════════════════════════════════════════════
print("\n[1] 소스 — 함수 경계 안에서 확인")
# ──────────────────────────────────────────────────────────────
validate_body = _body(RE_SRC, "_validate")
chk('str(val) == str(_c)' in validate_body,
    "★★★ ① _validate — choice 검증이 문자열/숫자 값을 함께 허용한다")

list_patterns_body = _body(RE_SRC, "list_patterns")
chk('"type", "choices", "default", "min", "max"' in list_patterns_body,
    "★★★ ① list_patterns — choices/type/min/max 를 화면에 그대로 내려준다")

chk("i.choices" in JS_SRC and "<select class=\"ag-in\"" in JS_SRC,
    "★★★ ① admin.js agPattern — choices 가 있으면 select 로 그린다")

posts_body = _body(RE_SRC, "_posts")
chk("math.ceil(L / s) + 1" in posts_body and "math.floor" not in posts_body,
    "★★★ ② _posts — floor 대신 ceil 을 쓴다(자기 docstring과 일치)")
#  ⚠ [라운드69] 정확한 버전 문자열을 그대로 박아 두면, 다음 라운드가 같은
#    컨벤션(파일 자신의 56행 주석: 레시피를 고치면 반드시 올린다)을 지켜
#    ENGINE_VER 를 또 올릴 때마다 이 회귀가 깨진다 — 실제로 69라운드에서 그랬다.
#    "58라운드 기준선(2026.07-1)보다 올라가 있다"만 확인한다(그 이후 몇 번을
#    더 올렸는지는 이 테스트의 관심사가 아니다).
chk('ENGINE_VER = "2026.07-1"' not in RE_SRC,
    "★ ② 수량식을 고쳤으니 ENGINE_VER 를 58라운드 기준선에서 올렸다(파일 자체 컨벤션)")

load_book_body = _body(APP_SRC, "_load_price_book")
chk('SELECT sku_key, unit, ref_price, price_src, price_date' in load_book_body,
    "★★★ ③ _load_price_book — unit 컬럼을 이제 읽는다")
chk('RE.SKU.get(r["sku_key"])' in load_book_body and
    '(r["unit"] or "").strip() != sku[2]' in load_book_body,
    "★★★ ③ _load_price_book — SKU 내장 단위와 다르면 그 행을 안 믿는다")
chk("COALESCE(is_demo,0)=0" in load_book_body,
    "★ ③(LOW-6) _load_price_book — 데모행을 걸러낸다")
chk("ORDER BY id" in load_book_body,
    "★ ③(LOW-6) _load_price_book — sku_key 중복 시 결정적 순서로 고른다")

#  ⚠ 삭제 이유를 남긴 주석 자체가 "SPREAD_MAX" 라는 이름을 언급한다(왜 지웠는지
#    설명하려면 옛 이름이 필요하다) — 그래서 이름의 완전한 부재가 아니라 정의문과
#    실제 비교식(if 문 전체)이 사라졌는가로 확인한다.
chk("SPREAD_MAX = 1.6" not in RE_SRC and
    "if low > 0 and high / low > SPREAD_MAX:" not in RE_SRC,
    "★ ④ _gate — 구조적으로 죽어 있던 SPREAD_MAX 게이트(정의+비교식)를 지웠다(비율은 늘 1.2)")

kpi_body = _body(APP_SRC, "api_ops_kpi_export")
chk("[_xl_safe(x) for x in r]" in kpi_body,
    "★ ⑤ api_ops_kpi_export — 형제 내보내기들과 같은 _xl_safe 정제를 거친다")

# ══════════════════════════════════════════════════════════════
print("\n[2] 실행 — 실제 엔진 호출로 ① choice 검증 + ② 수량식 재현")
# ──────────────────────────────────────────────────────────────
sys.path.insert(0, HERE)
import repair_estimate as RE  # noqa: E402

#  ⚠ 고치기 전 상태로 되돌리면 첫 estimate() 호출 자체가 EstimateError 를 던진다
#    (수정 전 회귀 그대로) — 크래시 대신 깔끔한 [FAIL] 로 보고되도록 감싼다.
try:
    # 화면이 실제로 보내는 그대로 — 문자열 "2"
    r_str2 = RE.estimate("guardrail_section", {"L": 9, "s": "2"})
    chk(r_str2["inputs"]["s"] == 2,
        "★★★ ① 문자열 \"2\" 가 더 이상 거부되지 않고 정수 2 로 정규화된다", r_str2["inputs"])

    r_str4 = RE.estimate("guardrail_section", {"L": 9, "s": "4"})
    chk(r_str4["inputs"]["s"] == 4,
        "★★ ① 문자열 \"4\" 도 정상 통과한다(선택지 둘 다 살아있다)", r_str4["inputs"])

    try:
        RE.estimate("guardrail_section", {"L": 9, "s": "3"})
        chk(False, "대조군 — 목록에 없는 값(\"3\")은 여전히 거부되어야 한다")
    except RE.EstimateError:
        chk(True, "★ 대조군 — 목록에 없는 값(\"3\")은 여전히 400 성격의 EstimateError")

    #  ⚠ 26.5% 과소산출의 실체 — 같은 L 에서 s=2(옛 기본값 s=4 대신 실제 좁은 간격)가
    #    s=4 보다 확실히 비싸게 나와야 한다(지주 더 많이 + m당 설치비도 더 비쌈).
    r_s2 = RE.estimate("guardrail_section", {"L": 9, "s": 2})
    r_s4 = RE.estimate("guardrail_section", {"L": 9, "s": 4})
    chk(r_s2["range_low"] > r_s4["range_low"] * 1.15,
        "★★★ ① s=2 견적이 옛 침묵 기본값(s=4)보다 실제로 유의미하게 높다"
        "(과소산출이 재현·해소됐다)", (r_s2["range_low"], r_s4["range_low"]))

    # ② _posts 자체 — L=9,s=2 는 배수가 아니므로 floor 였다면 1개 부족했다
    chk(RE._posts(9, 2) == 6, "★★★ ② _posts(9,2) == ceil(9/2)+1 == 6(고치기 전엔 5)", RE._posts(9, 2))
    chk(RE._posts(8, 4) == 3, "★ 대조군 — 배수인 경우(8,4)는 floor/ceil 결과가 같다(3)", RE._posts(8, 4))
    gr_post_qty = next(ln["qty"] for ln in r_s2["lines"] if ln["sku"] == "gr_post")
    chk(gr_post_qty == 6,
        "★★ ② 실제 견적 라인(gr_post 수량)에도 고쳐진 지주 수가 그대로 반영된다", gr_post_qty)
except Exception as e:                                              # noqa: BLE001
    chk(False, "★★★ ① 문자열 \"2\" 가 더 이상 거부되지 않고(엔진 직접호출 구간 전체)",
        f"{type(e).__name__} {e}")

# ══════════════════════════════════════════════════════════════
print("\n[3] 실행 — 실제 서버(test_client)로 ①(list_patterns) + ③ + ⑤ 재현")
# ──────────────────────────────────────────────────────────────
_tmpdir = tempfile.mkdtemp(prefix="round58_")
os.environ["BOKGUON_DB_PATH"] = os.path.join(_tmpdir, "verify.db")
os.environ["BOKGUON_UPLOAD_DIR"] = os.path.join(_tmpdir, "uploads")
os.environ["BOKGUON_ENV"] = "dev"
os.environ["BOKGUON_SEED_DEMO"] = "0"

try:
    import db as DB
    DB.init_db()
    import app as A

    client = A.app.test_client()
    _tok = "regress-round58-csrf-token"
    _hdr = {"X-CSRF-Token": _tok}

    conn = DB.get_db()
    conn.execute("INSERT INTO admins(username,password_hash,name,role) "
                 "VALUES('admin58','x','검증','중간관리자')")
    conn.commit(); conn.close()

    def _admin_sess(sess):
        sess["admin"] = "admin58"; sess["admin_role"] = "중간관리자"
        sess["admin_name"] = "검증"; sess["_csrf"] = _tok

    # ── ① /api/repair-patterns 가 choices/type 을 실제로 내려주는가 ──
    with client.session_transaction() as sess:
        _admin_sess(sess)
    rp = client.get("/api/repair-patterns")
    chk(rp.status_code == 200, "STEP /api/repair-patterns 정상 응답", rp.status_code)
    pats = {p["pattern"]: p for p in rp.get_json()["patterns"]} if rp.status_code == 200 else {}
    gr = pats.get("guardrail_section", {})
    s_field = next((i for i in gr.get("inputs", []) if i["key"] == "s"), {})
    chk(s_field.get("choices") == [2, 4] and s_field.get("type") == "choice",
        "★★★ ① API 가 choices=[2,4]·type=choice 를 그대로 내려준다"
        "(고치기 전엔 key·label만 내려가 화면이 선택지를 그릴 수 없었다)", s_field)

    # ── ① /api/repair-estimate 종단 — 문자열 "2" 로 실제 HTTP 재현 ──
    re1 = client.post("/api/repair-estimate",
                      json={"pattern": "guardrail_section", "inputs": {"L": 9, "s": "2"}},
                      headers=_hdr)
    chk(re1.status_code == 200,
        "★★★ ① HTTP 종단 — 문자열 \"2\" 로도 더 이상 400 이 아니다",
        (re1.status_code, re1.get_json()))
    if re1.status_code == 200:
        chk(re1.get_json()["estimate"]["inputs"]["s"] == 2,
            "★★ ① HTTP 응답의 s 도 2 로 정규화되어 있다(4 로 조용히 안 바뀜)",
            re1.get_json()["estimate"]["inputs"])

    # ── ③ _load_price_book 단위 불일치 재현 ──
    conn = DB.get_db()
    #  gr_panel 내장 단위="개" → 이 행은 unit='m' 으로 심어 불일치를 재현
    conn.execute(
        "INSERT INTO standard_boq(major,item,spec,unit,ref_price,sku_key,price_src,"
        "                         price_date,active,is_demo,created_at) "
        "VALUES(?,?,?,?,?,?,?,?,?,?,?)",
        ("가드레일", "가드레일 판(2W)", "4t×350×4330", "m", 999999999,
         "gr_panel", "검증58", "2026-08", 1, 0, DB.now()))
    #  gr_cap 내장 단위="개" → 이 행은 unit='개' 로 정확히 일치(대조군)
    conn.execute(
        "INSERT INTO standard_boq(major,item,spec,unit,ref_price,sku_key,price_src,"
        "                         price_date,active,is_demo,created_at) "
        "VALUES(?,?,?,?,?,?,?,?,?,?,?)",
        ("가드레일", "가드레일 캡", "Ø139.8", "개", 8800,
         "gr_cap", "검증58", "2026-08", 1, 0, DB.now()))
    #  gr_bolt_c 는 is_demo=1 로 심어 데모 필터 재현
    conn.execute(
        "INSERT INTO standard_boq(major,item,spec,unit,ref_price,sku_key,price_src,"
        "                         price_date,active,is_demo,created_at) "
        "VALUES(?,?,?,?,?,?,?,?,?,?,?)",
        ("가드레일", "취부볼트", "Ø19×180", "개", 5000,
         "gr_bolt_c", "검증58", "2026-08", 1, 1, DB.now()))
    conn.commit(); conn.close()

    conn = DB.get_db()
    book = A._load_price_book(conn)
    conn.close()
    chk("gr_panel" not in book,
        "★★★ ③ 단위가 어긋난 DB 단가행(개→m)은 엔진에 안 들어간다"
        "(고치기 전엔 unit 을 안 봐서 그대로 들어갔다)", book.get("gr_panel"))
    chk(book.get("gr_cap", {}).get("price") == 8800,
        "★ 대조군 — 단위가 맞는 행은 정상적으로 반영된다", book.get("gr_cap"))
    chk("gr_bolt_c" not in book,
        "★ ③(LOW-6) is_demo=1 행은 표준단가로 안 들어간다", book.get("gr_bolt_c"))

    # ③(LOW-6) 결정적 순서 — 같은 sku_key 두 행(단위 일치) 중 나중(id 큰) 값이 이긴다
    conn = DB.get_db()
    conn.execute(
        "INSERT INTO standard_boq(major,item,spec,unit,ref_price,sku_key,price_src,"
        "                         price_date,active,is_demo,created_at) "
        "VALUES(?,?,?,?,?,?,?,?,?,?,?)",
        ("교통안전", "시선유도표지 설치비", "흙속매설용", "개소", 1111,
         "delineator_ins", "검증58-구", "2026-01", 1, 0, DB.now()))
    conn.execute(
        "INSERT INTO standard_boq(major,item,spec,unit,ref_price,sku_key,price_src,"
        "                         price_date,active,is_demo,created_at) "
        "VALUES(?,?,?,?,?,?,?,?,?,?,?)",
        ("교통안전", "시선유도표지 설치비", "흙속매설용", "개소", 2222,
         "delineator_ins", "검증58-신", "2026-08", 1, 0, DB.now()))
    conn.commit(); conn.close()
    conn = DB.get_db()
    book2 = A._load_price_book(conn)
    conn.close()
    chk(book2.get("delineator_ins", {}).get("price") == 2222,
        "★ ③(LOW-6) sku_key 중복 시 ORDER BY id 로 나중 값이 결정적으로 이긴다", book2.get("delineator_ins"))

    # ── ⑤ 엑셀 내보내기가 정상 동작(회귀 아님)하는지 스모크 ──
    rk = client.get("/api/ops-kpi/export")
    chk(rk.status_code == 200, "★ ⑤ /api/ops-kpi/export — _xl_safe 적용 후에도 정상 200", rk.status_code)
    if rk.status_code == 200:
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(rk.data))
        ws = wb.active
        chk(ws["A3"].value == "구분",
            "★ ⑤ 정제 후에도 헤더 셀 내용이 그대로다(문자열 아닌 값은 안 건드림)", ws["A3"].value)

except Exception as e:                                              # noqa: BLE001
    chk(False, "★★★ 라운드58 실측 준비/실행", f"{type(e).__name__} {e}")
finally:
    os.environ.pop("BOKGUON_DB_PATH", None)
    os.environ.pop("BOKGUON_UPLOAD_DIR", None)
    os.environ.pop("BOKGUON_ENV", None)
    os.environ.pop("BOKGUON_SEED_DEMO", None)
    import shutil
    shutil.rmtree(_tmpdir, ignore_errors=True)

print("\n" + "═" * 72)
print(f"  통과 {len(PASS)} · 실패 {len(FAIL)}")
if FAIL:
    print("\n  실패 항목:")
    for nm, dt in FAIL:
        print(f"    · {nm}\n        {dt}")
print("═" * 72)
sys.exit(1 if FAIL else 0)
