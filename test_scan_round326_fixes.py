# -*- coding: utf-8 -*-
"""스캔 326라운드 회귀 — 감사로그 화면·엑셀(/api/activity, /api/activity/export)
   이 관리자 계정 관리(생성/등급변경/이름변경/비밀번호재발급/삭제)의 아이디·
   이름·등급을 등급 무관하게 노출하던 결함[중간, 접근통제 우회], 이 세션
   사상 첫 정밀조사한 정밀조사 Agent가 발견(app.py 480개 함수/라우트를
   "가장 가까운 라운드 주석까지의 최소거리"로 감사밀도를 계산해 재정렬,
   최저밀도(317줄) 1위인 api_activity 를 선택 — 320~325라운드가 이어온
   app.py 함수단위 재구성 기법의 3회차).

   `GET /api/admins`(관리자 계정 목록)는 `@require_role("최고관리자")`로
   엄격히 제한되는데, `GET /api/activity`·`GET /api/activity/export`는
   등급 검사가 없는 `@require_admin`만 요구한다("하위관리자도 감사로그는
   본다"는 의도된 설계 — 이 자체는 정당하고 안 바꾼다). 그런데 관리자
   계정 생성(app.py:1648)·등급변경(1680)·이름변경(1682)·비밀번호재발급
   (1692)·삭제(1719)의 DB.log() detail 에 대상 계정의 아이디·이름·등급이
   그대로 찍혀, 하위 등급 관리자가 `/api/admins`에서는 403 을 받으면서도
   감사로그 화면·엑셀로 똑같은 정보를 우회해서 볼 수 있었다(실측:
   중간관리자 세션으로 재현). 행 자체(시각·행위자·행위명)는 그대로 두고
   이 5개 행위의 detail 만 최고관리자 미만에게 가리는 `_mask_admin_activity()`
   공유 헬퍼를 신설해 화면·엑셀 양쪽에 적용.
"""
import io
import os
import sys
import tempfile

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)

_scratch = tempfile.mkdtemp(prefix="bokguon_r326_")
os.environ["BOKGUON_DB_PATH"] = os.path.join(_scratch, "verify.db")
os.environ["BOKGUON_ENV"] = "dev"

PASS = FAIL = 0
FAILED = []


def chk(cond, msg, detail=None):
    global PASS, FAIL
    if cond:
        PASS += 1
        print(f"  [OK]   {msg}")
    else:
        FAIL += 1
        FAILED.append(msg)
        print(f"  [FAIL] {msg}" + (f"   {detail}" if detail else ""))


print("=" * 74)
print("  스캔 326라운드 회귀 — 감사로그가 관리자 계정 관리 정보를 등급 무관 노출")
print("=" * 74)

import db as DB  # noqa: E402
DB.init_db(seed_demo=True)
import app as A  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

app_client = A.app.test_client()
_tok = "r326-admin-csrf"

print("\n[1] 준비 — 최고관리자로 새 관리자 계정을 생성해 실제 감사로그를 남긴다")
with app_client.session_transaction() as sess:
    sess["admin"] = "admin"
    sess["admin_role"] = "최고관리자"
    sess["_csrf"] = _tok
_hdr = {"X-CSRF-Token": _tok}
r_create = app_client.post("/api/admins", headers=_hdr,
                            json={"username": "r326target", "name": "몰래볼대상", "role": "하위관리자"})
chk(r_create.status_code == 200 and r_create.get_json().get("ok") is True,
    "  관리자 계정 생성 자체는 정상 200/ok", (r_create.status_code, r_create.get_data(as_text=True)))

#  대조군 — 관리자 계정과 무관한 일반 감사로그(하위관리자도 봐야 정상)
conn = DB.get_db()
DB.log(conn, "테스트관리자", "사건 등록(직접등록)", "26-99999 · R326-대조군")
conn.commit()

print("\n[2] 실행 — 중간관리자(최고관리자 미만) 세션으로 감사로그 화면을 본다")
with app_client.session_transaction() as sess:
    sess["admin"] = "staff"
    sess["admin_role"] = "중간관리자"
    sess["_csrf"] = _tok
r_view = app_client.get("/api/activity?limit=200", headers=_hdr)
chk(r_view.status_code == 200, "  감사로그 화면 자체는 정상 200(하위 등급도 조회 가능 — 원래 의도)",
    r_view.status_code)
items = r_view.get_json().get("items", [])
create_rows = [it for it in items if it.get("action") == "관리자 계정 생성"]
control_rows = [it for it in items if it.get("action") == "사건 등록(직접등록)"]

chk(len(create_rows) >= 1, "  관리자 계정 생성 로그 행 자체는 여전히 보인다(행 삭제 아님)", len(create_rows))
if create_rows:
    chk("r326target" not in create_rows[0]["detail"] and "몰래볼대상" not in create_rows[0]["detail"],
        "★★★★★ 중간관리자에게는 관리자 계정 생성의 아이디·이름이 감춰진다"
        "(고치기 전엔 detail 에 그대로 노출됐다)", create_rows[0]["detail"])
chk(len(control_rows) == 1 and "R326-대조군" in control_rows[0]["detail"],
    "  관리자 계정과 무관한 일반 로그는 중간관리자에게도 그대로 보인다(회귀 없음 — 의도된 설계 유지)",
    control_rows[0] if control_rows else None)

print("\n[3] 실행 — 같은 조건으로 엑셀 내보내기도 가려진다")
r_export = app_client.get("/api/activity/export?limit=200", headers=_hdr)
chk(r_export.status_code == 200, "  엑셀 내보내기 자체는 정상 200", r_export.status_code)
wb = load_workbook(io.BytesIO(r_export.data))
ws = wb.active
xl_rows = list(ws.iter_rows(min_row=4, values_only=True))  # 1~3행은 제목·요약·헤더
xl_create = [r for r in xl_rows if r[2] == "관리자 계정 생성"]
xl_control = [r for r in xl_rows if r[2] == "사건 등록(직접등록)"]
chk(len(xl_create) >= 1 and "r326target" not in (xl_create[0][3] or ""),
    "★★★★ 엑셀에서도 관리자 계정 생성의 아이디가 가려진다", xl_create[0] if xl_create else None)
chk(len(xl_control) == 1 and "R326-대조군" in (xl_control[0][3] or ""),
    "  엑셀의 일반 로그는 그대로 보인다(회귀 없음)", xl_control[0] if xl_control else None)

print("\n[4] 실행 — 최고관리자 세션으로는 여전히 전부 그대로 보인다(회귀 없음)")
with app_client.session_transaction() as sess:
    sess["admin"] = "admin"
    sess["admin_role"] = "최고관리자"
    sess["_csrf"] = _tok
r_super = app_client.get("/api/activity?limit=200", headers=_hdr)
super_create = [it for it in r_super.get_json().get("items", []) if it.get("action") == "관리자 계정 생성"]
chk(len(super_create) >= 1 and "r326target" in super_create[0]["detail"],
    "  최고관리자에게는 여전히 아이디·이름이 그대로 보인다(과잉차단 아님)",
    super_create[0] if super_create else None)

print("\n" + "=" * 74)
print(f"  통과 {PASS} · 실패 {FAIL}")
if FAILED:
    print("\n  실패 항목:")
    for m in FAILED:
        print(f"    · {m}")
print("=" * 74)

sys.exit(1 if FAIL else 0)
