# -*- coding: utf-8 -*-
"""스캔 409라운드 회귀 — api_contact_import()(엑셀 일괄등록, app.py:8632)의
   MAX_IMPORT_ROWS(2000) 상한이 시트 XML의 <dimension> 태그(파일이 스스로
   선언하는 메타데이터)에 의존해, 그 태그가 없는 파일로 완전히 우회되던 결함.

  배경: 409라운드 조사 Agent가 발견 — n_rows = max((ws.max_row or 1) - 1, 0)
  (app.py:8663)는 <dimension> 태그가 없으면 ws.max_row 가 None 이 되어
  n_rows=0 으로 계산돼 상한검사를 무조건 통과한다. 그런데 실제 처리 루프
  (`for row in ws.iter_rows(min_row=2):`, 8677행)도 max_row 인자를 안 줘서
  똑같이 None 인 self.max_row 에 기대므로(openpyxl 내부 구현) 내부 상한
  없이 파일에 실제로 있는 행을 전부 스트리밍한다 — 가드와 처리루프가 같은
  선언값을 "함께 신뢰"하는 게 아니라, 선언이 사라지면 가드는 0으로,
  처리는 무제한으로 서로 다르게 반응한다.

  형제함수(api_partner_bulk_credentials)는 실제 파싱된 파이썬 리스트를
  직접 슬라이싱해 이 문제가 없고, 사진업로드 경로(11360-64행)는 이미
  "실제 처리 개수로 커밋 직전에 다시 검사한다" 원칙을 쓰고 있다 — 같은
  원칙을 여기로 그대로 가져와 고쳤다.

  고침: 루프 안에서 실제로 센 행수(_real_rows)가 MAX_IMPORT_ROWS 를
  넘으면 _row_limit_hit 를 세우고 즉시 break, 루프 뒤에서 rollback+400.
  기존 사전검사(n_rows, <dimension> 태그 기반)는 정상 파일에 대한
  빠른 거부용 1차 방어로 그대로 둔다(정직한 초과 파일은 트랜잭션을
  아예 안 열고 거부 — round70의 기존 회귀테스트가 이미 검증).

      python test_scan_round409_fixes.py
"""
import io
import os
import re
import sys
import tempfile
import zipfile

HERE = os.path.dirname(os.path.abspath(__file__))
PASS, FAIL = [], []


def chk(cond, name, detail=""):
    (PASS if cond else FAIL).append((name, detail))
    print(("  [OK]   " if cond else "  [FAIL] ") + name + ("" if cond else f"   {detail}"))


def read(rel):
    return io.open(os.path.join(HERE, rel), encoding="utf-8").read()


APP_SRC = read("app.py")

#  고친 뒤(현재) 실제 처리루프의 시작부와 되돌릴 옛(라운드70 시절) 시작부.
NEW_LOOP_START = (
    '        try:\n'
    '            #  ⚠⚠ [라운드409]'
)
OLD_LOOP_START = (
    '        try:\n'
    '            for row in ws.iter_rows(min_row=2):\n'
    '                def cell(k):'
)
NEW_ROW_LIMIT_CHECK = (
    '            if _row_limit_hit:\n'
    '                conn.rollback()\n'
    '                return jsonify({"ok": False,\n'
    '                                "error": f"한 번에 최대 {MAX_IMPORT_ROWS}행까지 처리할 수 있습니다. "\n'
    '                                         f"파일을 나눠서 올려 주세요."}), 400\n'
    '            DB.log(conn,'
)
OLD_NO_ROW_LIMIT_CHECK = '            DB.log(conn,'

print("=" * 72)
print("  스캔 409라운드 회귀 — 엑셀 일괄등록 dimension태그 부재로 행수상한 우회")
print("=" * 72)

print("\n[1] 소스 확인")
chk("_real_rows = 0" in APP_SRC and "_row_limit_hit" in APP_SRC,
    "★★★★★ api_contact_import() 가 실제로 센 행수로 상한을 다시 강제한다")
chk(NEW_ROW_LIMIT_CHECK in APP_SRC,
    "  루프 뒤 rollback+400 처리가 정확한 위치에 있다")


def strip_dimension_tag(xlsx_bytes):
    """openpyxl 이 만든 xlsx 의 시트 XML에서 <dimension .../> 태그를 제거한다
       (409라운드 조사 Agent가 실측 확인한 우회 기법과 동일)."""
    zin = zipfile.ZipFile(io.BytesIO(xlsx_bytes), "r")
    sheet_name = [n for n in zin.namelist()
                  if n.startswith("xl/worksheets/sheet") and n.endswith(".xml")][0]
    sheet_xml = zin.read(sheet_name).decode("utf-8")
    new_xml = re.sub(r'<dimension ref="[^"]+"/>', "", sheet_xml)
    if new_xml == sheet_xml:
        raise AssertionError("dimension 태그를 못 찾음 — openpyxl 출력 형식이 바뀌었을 수 있음")
    out_buf = io.BytesIO()
    zout = zipfile.ZipFile(out_buf, "w", zipfile.ZIP_DEFLATED)
    for item in zin.infolist():
        data = zin.read(item.filename)
        if item.filename == sheet_name:
            data = new_xml.encode("utf-8")
        zout.writestr(item, data)
    zout.close()
    zin.close()
    return out_buf.getvalue()


def make_xlsx(rows):
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.append(["업체명", "이메일"])
    for r in rows:
        ws.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


print("\n[2] 실행 — 실제 서버 호출로 재현")
_tmpdir = tempfile.mkdtemp(prefix="round409_")
DBPATH = os.path.join(_tmpdir, "verify.db")
os.environ["BOKGUON_DB_PATH"] = DBPATH
os.environ["BOKGUON_UPLOAD_DIR"] = os.path.join(_tmpdir, "uploads")
os.environ["BOKGUON_ENV"] = "dev"
os.environ["BOKGUON_SEED_DEMO"] = "0"
sys.path.insert(0, HERE)

N_MALICIOUS_ROWS = 2500
TOK = "regress-round409-csrf-token"

try:
    import db as DB
    DB.init_db()
    import app as A

    conn = DB.get_db()
    conn.execute("INSERT INTO admins(username,password_hash,name,role) "
                 "VALUES('admin','x','검증','최고관리자')")
    #  위조파일의 각 행이 실제로 UPDATE 를 만들어내는지(단순 통과가 아니라
    #  진짜 처리됨)까지 실측하기 위해, 행 수만큼 이름이 일치하는 업체를 미리 심는다.
    for i in range(N_MALICIOUS_ROWS):
        conn.execute("INSERT INTO partners(name, email) VALUES(?, ?)",
                     (f"R409업체{i}", f"old{i}@x.com"))
    conn.commit(); conn.close()

    client = A.app.test_client()
    with client.session_transaction() as sess:
        sess["admin"] = "admin"; sess["admin_role"] = "최고관리자"; sess["admin_name"] = "검증"
        sess["_csrf"] = TOK
    hdr = {"X-CSRF-Token": TOK}

    malicious_bytes = strip_dimension_tag(
        make_xlsx([[f"R409업체{i}", f"new{i}@x.com"] for i in range(N_MALICIOUS_ROWS)]))

    def count_updated():
        c = DB.get_db()
        n = c.execute("SELECT COUNT(*) c FROM partners WHERE email LIKE 'new%@x.com'").fetchone()["c"]
        c.close()
        return n

    r1 = client.post("/api/contacts/import",
                      data={"file": (io.BytesIO(malicious_bytes), "evil.xlsx")},
                      headers=hdr, content_type="multipart/form-data")
    j1 = r1.get_json() or {}
    chk(r1.status_code == 400 and not j1.get("ok") and "2000행" in (j1.get("error") or ""),
        "★★★★★ dimension 태그가 없는 2500행 위조파일은 지금은 400으로 거부된다"
        "(고치기 전엔 가드가 0행으로 오판해 그대로 통과시켰다)", j1)
    chk(count_updated() == 0,
        "★★★★★ 거부된 요청은 실제로 롤백되어 partners 테이블에 반영된 게 하나도 없다",
        count_updated())

    print("\n[3] 되돌리기 — 옛 코드(루프 안 실제카운트 없음)로 되돌리면 재현되는가")
    #  ⚠⚠ [라운드612] round412·414가 확립한 안전패턴으로 이관 — 원래는 되돌린 전체
    #    app.py 소스를 실제 on-disk APP_PATH 에 잠깐 써넣고 subprocess로 재현한 뒤
    #    finally에서 복원했다(round409~515대 옛 세대 기법). 이 저장소는 git이 없어
    #    동시세션이 상시전제라, write→subprocess.run→write 사이의 창을 다른 세션이
    #    관측하면 충돌한다(round590이 실제로 겪음). 함수 본문 문자열만 옛 조건으로
    #    되돌려 exec(compile(...))로 네임스페이스 사본에 재구성하고, A.app.view_
    #    functions["api_contact_import"] 를 잠깐 치환해 같은 test_client() 로 직접
    #    호출한다(실제 파일은 전혀 안 건드림) — 증명하려는 인과관계(옛 코드는 위조
    #    파일을 무제한 통과시킨다)는 그대로다.
    chk(NEW_LOOP_START in APP_SRC and OLD_LOOP_START not in APP_SRC,
        "  되돌리기 준비 — 현재 소스가 고친 형태임을 확인")

    _fn_start = APP_SRC.index("\ndef api_contact_import(")
    _fn_end = APP_SRC.index("\n\n\n", _fn_start)
    _cur_fn_src = APP_SRC[_fn_start:_fn_end]
    chk(NEW_LOOP_START in _cur_fn_src and NEW_ROW_LIMIT_CHECK in _cur_fn_src,
        "  되돌리기 준비 — 함수 본문 안에 두 조각이 모두 있다")

    i0 = _cur_fn_src.index(NEW_LOOP_START)
    i1 = _cur_fn_src.index("                def cell(k):")
    reverted_fn = (_cur_fn_src[:i0] + OLD_LOOP_START
                   + _cur_fn_src[i1 + len("                def cell(k):"):])
    reverted_fn = reverted_fn.replace(NEW_ROW_LIMIT_CHECK, OLD_NO_ROW_LIMIT_CHECK)
    reverted_fn = reverted_fn.replace(
        "def api_contact_import():", "def _api_contact_import_pre409():", 1)
    chk("_row_limit_hit" not in reverted_fn and "for row in ws.iter_rows(min_row=2):" in reverted_fn
        and "_api_contact_import_pre409" in reverted_fn,
        "  되돌리기 준비 — 되돌린 함수본문에 새 코드 흔적이 안 남는다(실제 파일은 안 건드림)")

    _ns = dict(vars(A))
    exec(compile(reverted_fn, "<round409_pre409_contact_import>", "exec"), _ns)  # noqa: S102 -- 우리 소스에서 그대로 추출, 디스크 미기록
    old_contact_import_fn = A.require_admin(_ns["_api_contact_import_pre409"])

    _orig_view = A.app.view_functions["api_contact_import"]
    A.app.view_functions["api_contact_import"] = old_contact_import_fn
    try:
        r_old = client.post("/api/contacts/import",
                             data={"file": (io.BytesIO(malicious_bytes), "evil.xlsx")},
                             headers=hdr, content_type="multipart/form-data")
    finally:
        A.app.view_functions["api_contact_import"] = _orig_view
    j_old = r_old.get_json() or {}
    chk(r_old.status_code == 200 and j_old.get("updated") == N_MALICIOUS_ROWS,
        "★★★★★ 옛 코드(재구성, 실제 파일은 안 건드림)로 되돌리면 같은 위조파일이 200으로 "
        "통과해 전부 갱신됨(인과관계 직접증명)", j_old)
    chk(count_updated() == N_MALICIOUS_ROWS,
        "  실제로 partners 테이블에도 반영까지 됐다(트랜잭션이 커밋까지 감, 롤백 아님)",
        count_updated())

    print("\n[4] 정상 경로 — 대조군(정상 소량 파일)은 회귀없이 그대로 동작")
    ok_bytes = make_xlsx([["R409업체0", "final0@x.com"]])
    r4 = client.post("/api/contacts/import",
                      data={"file": (io.BytesIO(ok_bytes), "ok.xlsx")},
                      headers=hdr, content_type="multipart/form-data")
    j4 = r4.get_json() or {}
    chk(r4.status_code == 200 and j4.get("ok") and j4.get("updated") == 1,
        "  정상 소량 파일(1행)은 지금도 정상 처리된다(회귀없음)", j4)
except Exception as e:                                              # noqa: BLE001
    chk(False, "★★★★★ 준비/실행", f"{type(e).__name__} {e}")
#  [라운드612] 실제 app.py를 이제 전혀 안 쓰므로(위 [3]을 exec 재구성으로 이관) 이
#  안전망(디스크 복원 finally)도 함께 걷어냈다 — 더 이상 쓰는 곳이 없다.

print("\n" + "=" * 72)
print(f"  통과 {len(PASS)} · 실패 {len(FAIL)}")
if FAIL:
    print("\n  실패 항목:")
    for nm, dt in FAIL:
        print(f"    · {nm}\n        {dt}")
print("=" * 72)
sys.exit(1 if FAIL else 0)
