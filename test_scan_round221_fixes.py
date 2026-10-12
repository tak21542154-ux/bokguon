# -*- coding: utf-8 -*-
"""스캔 221라운드 회귀 — 두 결함.

  [1] api_export_xlsx() 가 모듈 전역 헬퍼 _xl_safe() 를 쓰지 않고 독자적인
      로컬 _safe()(ILLEGAL_CHARACTERS_RE 기반)를 따로 재구현하고 있었다.
      _xl_safe() 자신의 문서가 "내보내기 어디서든 같은 규칙을 쓴다"고
      명시하는데 이 함수만 어겼다 — 방어를 두 벌로 유지하는 이 클래스의
      결함은 이 코드베이스에서 이미 실제 사고(보험사 명세서에만 규칙이
      빠졌던 사건)를 낸 전례가 있다. 실행 검증한 유일한 행동 차이는 \\r
      (캐리지리턴) 처리 — 로컬 _safe 는 보존, _xl_safe 는 제거. 이 섹션
      (applications·leads)은 비인증 공개폼(apply_partner.html 등)의
      자유텍스트가 그대로 들어오는 경로라 방어가 흩어져 있으면 안 된다.
  [2] static/style.css 의 body font-family 가 OS별 폴백을 손으로만 일부
      (Windows) 적어 둬, fonts.css 가 이미 정의해 둔 전체 대체 사슬
      (macOS: Apple SD Gothic Neo, Android: Noto Sans KR 포함, --font-sans)
      을 랜딩·에러·약관 3개 페이지만 쓰고 나머지 콘솔·포털 14개 템플릿은
      못 받고 있었다.

      python test_scan_round221_fixes.py
"""
import io
import os
import shutil
import sys
import tempfile

HERE = os.path.dirname(os.path.abspath(__file__))
PASS, FAIL = [], []


def chk(cond, name, detail=""):
    (PASS if cond else FAIL).append((name, detail))
    print(("  [OK]   " if cond else "  [FAIL] ") + name + ("" if cond else f"   {detail}"))


def read(rel):
    return io.open(os.path.join(HERE, rel), encoding="utf-8").read()


def _py_fn_body(src, fn_name):
    start = src.index(f"\ndef {fn_name}(")
    end = src.find("\ndef ", start + 1)
    return src[start:end if end != -1 else None]


APP_SRC = read("app.py")
STYLE_SRC = read("static/style.css")
FONTS_SRC = read("static/fonts.css")

print("═" * 72)
print("  스캔 221라운드 회귀 — 엑셀정제 헬퍼 이중구현 · body 폰트폴백 누락")
print("═" * 72)

# ══════════════════════════════════════════════════════════════
print("\n[1] 소스 확인 — api_export_xlsx 함수 경계 안")
# ──────────────────────────────────────────────────────────────
export_body = _py_fn_body(APP_SRC, "api_export_xlsx")
chk("def _safe(" not in export_body,
    "★★★ api_export_xlsx 가 더 이상 로컬 _safe() 를 독자 재구현하지 않는다")
chk("ILLEGAL_CHARACTERS_RE" not in export_body,
    "STEP 그 로컬 구현이 쓰던 import 도 함께 정리됐다(과잉삭제 아님 — 아래서 "
    "_xl_safe 로 실제 동작함을 실행으로 확인)")
chk("return _xl_safe(v)" in export_body and export_body.count("_xl_safe(v)") == 2,
    "★★★ 대신 모듈 공유 헬퍼 _xl_safe() 를 두 분기(날짜 파싱 실패 시·일반)"
    "모두에서 호출한다", export_body.count("_xl_safe(v)"))

chk("font-family:var(--font-sans)" in STYLE_SRC,
    "★★★★ style.css 의 body font-family 가 fonts.css 의 --font-sans 를 "
    "그대로 쓴다(예전엔 맑은 고딕·Malgun Gothic 만 손으로 적어 macOS·Android "
    "폴백이 여기서만 빠졌다)")
chk("--font-sans:" in FONTS_SRC and "Apple SD Gothic Neo" in FONTS_SRC
    and "Noto Sans KR" in FONTS_SRC,
    "STEP fonts.css 쪽 --font-sans 정의 자체는 그대로다(참고용 확인)")

# ══════════════════════════════════════════════════════════════
print("\n[2] 실행 — 실서버 경로(Flask test_client + 임시 DB)로 실제 재현")
# ──────────────────────────────────────────────────────────────
_tmpdir = tempfile.mkdtemp(prefix="round221_")
os.environ["BOKGUON_DB_PATH"] = os.path.join(_tmpdir, "verify.db")
os.environ["BOKGUON_ENV"] = "dev"
os.environ["BOKGUON_SEED_DEMO"] = "0"
sys.path.insert(0, HERE)
try:
    import db as DB
    DB.init_db()
    import app as A
    from openpyxl import load_workbook

    #  비인증 공개폼(apply_partner.html)이 실제로 받는 자유텍스트를 흉내낸다 —
    #  수식주입 시도(company)와, 캐리지리턴이 섞인 메모(note, 워드 붙여넣기
    #  등에서 흔함) 둘 다 같은 행에 넣는다.
    MALICIOUS_COMPANY = "=합계요청업체"
    NOTE_WITH_CR = "1차검토완료\r2차검토대기"

    conn = DB.get_db()
    conn.execute("INSERT INTO admins(username,password_hash,name,role) "
                 "VALUES('r221admin','x','검증','중간관리자')")
    conn.execute(
        "INSERT INTO partner_applications(company,status,note) VALUES(?,?,?)",
        (MALICIOUS_COMPANY, "신규", NOTE_WITH_CR))
    conn.commit()
    conn.close()

    admin = A.app.test_client()
    with admin.session_transaction() as sess:
        sess["admin"] = "r221admin"; sess["admin_role"] = "중간관리자"; sess["admin_name"] = "검증"

    r = admin.get("/api/export/applications")
    _d = "" if r.status_code == 200 else r.get_data(as_text=True)[:300]
    chk(r.status_code == 200,
        "STEP 입점신청 엑셀 내보내기 자체는 정상 처리된다(과잉차단 아님)",
        f"{r.status_code} {_d}")

    if r.status_code == 200:
        wb = load_workbook(io.BytesIO(r.get_data()))
        ws = wb.active
        headers = [c.value for c in ws[1]]
        row = [c.value for c in ws[2]]
        rec = dict(zip(headers, row))

        chk(rec.get("업체명") == "'" + MALICIOUS_COMPANY,
            "★★★ 수식주입 시도 문자열은 여전히 앞에 ' 가 붙어 무력화된다"
            "(_xl_safe 전환 후에도 기존 방어가 그대로 살아있다)",
            rec.get("업체명"))
        chk(rec.get("심사메모") == "1차검토완료2차검토대기",
            "★★★★ 캐리지리턴(\\r)이 이제 _xl_safe 규칙대로 완전히 제거된다"
            "(고치기 전 로컬 _safe 는 \\r 을 그대로 남겨 두 헬퍼가 서로 다른 "
            "결과를 냈다)",
            repr(rec.get("심사메모")))

except Exception as e:                                              # noqa: BLE001
    chk(False, "★★★ 라운드221 서버측 준비/실행", f"{type(e).__name__} {e}")
finally:
    os.environ.pop("BOKGUON_DB_PATH", None)
    os.environ.pop("BOKGUON_ENV", None)
    os.environ.pop("BOKGUON_SEED_DEMO", None)
    shutil.rmtree(_tmpdir, ignore_errors=True)

print("\n" + "═" * 72)
print(f"  통과 {len(PASS)} · 실패 {len(FAIL)}")
if FAIL:
    print("\n  실패 항목:")
    for nm, dt in FAIL:
        print(f"    · {nm}\n        {dt}")
print("═" * 72)
sys.exit(1 if FAIL else 0)
