# -*- coding: utf-8 -*-
"""스캔 142라운드 회귀 — 2건.

  [1] 엑셀 내보내기(하자AS·입점신청·도입문의)가 화면 필터를 통째로 무시했다.
      cases/partners/settlements 는 이미 화면과 같은 조건으로 엑셀을 내는데
      (122·104라운드 등), 이 세 섹션은 EXPORT_SECTIONS 의 SQL에 {WHERE} 자리가
      없어 목록 API(search_filters 기반)가 받는 status·category·q·기간 조건을
      엑셀만 무시하고 항상 전체(비-데모)를 냈다.
  [2] manage.py prune_personal 의 '종료(취소·미수선종결·유찰)' 버킷(n4)이
      cases.victim_name/phone 만 지우고, 형제 버킷(완료 사건, n3)이 이미 하는
      후속 파기 2가지(notifications 본문 redaction·track_token 무효화)를
      안 했다 — 90일 지난 유찰 사건의 진행조회 링크가 파기 로그가 남은
      뒤에도 그대로 열려 있었다(개인정보 보호법 §21 취지 위반).

      python test_scan_round142_fixes.py
"""
import io
import os
import sys
import tempfile
from datetime import datetime, timedelta

HERE = os.path.dirname(os.path.abspath(__file__))
PASS, FAIL = [], []


def chk(cond, name, detail=""):
    (PASS if cond else FAIL).append((name, detail))
    print(("  [OK]   " if cond else "  [FAIL] ") + name + ("" if cond else f"   {detail}"))


def read(rel):
    return io.open(os.path.join(HERE, rel), encoding="utf-8").read()


APP_SRC = read("app.py")
MANAGE_SRC = read("manage.py")
JS_SRC = read("static/admin.js")

print("=" * 72)
print("  스캔 142라운드 회귀 — 엑셀 내보내기 필터 무시(3섹션) · 개인정보 파기 후속처리 누락")
print("=" * 72)

# ============================================================
print("\n[1] 소스 확인")
# ------------------------------------------------------------
chk("def _lead_filter(args):" in APP_SRC and "def _application_filter(args):" in APP_SRC
    and "def _defect_filter(args):" in APP_SRC,
    "★★★ 도입문의·입점신청·하자AS 전용 엑셀 필터 함수 3개가 신설됐다")
chk('"leads": _lead_filter' in APP_SRC and '"applications": _application_filter' in APP_SRC
    and '"defects": _defect_filter' in APP_SRC,
    "★★★ _EXPORT_FILTERS 에 3개 필터가 전부 등록됐다")
chk('FROM insurer_leads {WHERE}' in APP_SRC, "★ leads SQL 이 {WHERE} 자리를 갖는다")
chk('FROM partner_applications {WHERE}' in APP_SRC, "★ applications SQL 이 {WHERE} 자리를 갖는다")
chk('FROM defects {WHERE}\n' in APP_SRC, "★ defects SQL 이 {WHERE} 자리를 갖는다")

chk("(a === 'defects') ? defectQS().toString()" in JS_SRC
    and "(a === 'leads') ? leadQS().toString()" in JS_SRC
    and "(a === 'applications') ? appQS().toString()" in JS_SRC,
    "★★★ admin.js exportXlsx 가 하자·도입문의·입점신청 화면필터를 실어 보낸다")

chk("종료(취소·미수선·유찰) 피해자 안내 알림" in MANAGE_SRC,
    "★★★ prune_personal 이 종료버킷(n4)의 알림 본문도 파기한다")
chk("종료(취소·미수선·유찰) 진행조회 링크" in MANAGE_SRC,
    "★★★ prune_personal 이 종료버킷(n4)의 진행조회 토큰도 무효화한다")
chk("n1 + n2 + n3 + n3b + n3c + n3d + n4 + n4b + n4c" in MANAGE_SRC,
    "★★★ total 집계가 새 카운터 n4b·n4c 를 포함한다(109라운드 교훈 재적용)")

# ============================================================
print("\n[2] 실행 준비 — 실제 app.py(전체 임포트) + 어드민 세션")
# ------------------------------------------------------------
_tmpdir = tempfile.mkdtemp(prefix="round142_")
os.environ["BOKGUON_DB_PATH"] = os.path.join(_tmpdir, "verify.db")
os.environ["BOKGUON_UPLOAD_DIR"] = os.path.join(_tmpdir, "uploads")
os.environ["BOKGUON_ENV"] = "dev"
os.environ["BOKGUON_SEED_DEMO"] = "0"
sys.path.insert(0, HERE)

try:
    import db as DB
    DB.init_db()
    import app as A
    from openpyxl import load_workbook

    conn = DB.get_db()
    conn.execute("INSERT INTO admins(username,password_hash,name,role) "
                 "VALUES('admin','x','검증관리자','최고관리자')")
    conn.commit()
    admin = A.app.test_client()
    tok = "r142-admin-csrf"
    with admin.session_transaction() as sess:
        sess["admin"] = "admin"; sess["admin_role"] = "최고관리자"; sess["admin_name"] = "검증관리자"
        sess["_csrf"] = tok

    def _xlsx_rows(resp, id_col):
        wb = load_workbook(io.BytesIO(resp.data))
        ws = wb.active
        headers = [c.value for c in ws[1]]
        ci = headers.index(id_col)
        return [r[ci] for r in ws.iter_rows(min_row=2, values_only=True)]

    # ── [1] 하자AS: 상태='접수' 2건 + 다른 상태 3건, 화면 필터=접수 ──────
    for i, st in enumerate(["접수", "접수", "확인중", "완료", "반려"]):
        conn.execute(
            "INSERT INTO defects(defect_no,title,status,is_demo) "
            f"VALUES('D26-R142{i}','R142하자{i}',?,0)", (st,))
    conn.commit()
    r_list = admin.get("/api/defects?status=접수")
    chk(r_list.status_code == 200, "STEP 하자 목록(status=접수) 200")
    n_list = len(r_list.get_json())
    chk(n_list == 2, "★ (사전조건) 목록이 status=접수 로 2건만 돌려준다", n_list)
    r_xl = admin.get("/api/export/defects?status=접수")
    chk(r_xl.status_code == 200, "STEP 하자 엑셀 내보내기 200")
    titles = _xlsx_rows(r_xl, "제목")
    chk(len(titles) == 2,
        "★★★ 엑셀도 목록과 같은 status=접수 필터를 받아 2건만 담는다"
        "(고치기 전엔 필터 무시하고 5건 전체가 나갔다)", titles)

    # ── [2] 입점신청: 상태='신규' 2건 + 다른 상태 2건, 화면 필터=신규 ────
    for i, st in enumerate(["신규", "신규", "승인", "반려"]):
        conn.execute(
            "INSERT INTO partner_applications(company,status,is_demo) "
            f"VALUES('R142신청업체{i}',?,0)", (st,))
    conn.commit()
    r_list = admin.get("/api/applications?status=신규")
    n_list = len(r_list.get_json())
    chk(n_list == 2, "★ (사전조건) 입점신청 목록이 status=신규 로 2건만 돌려준다", n_list)
    r_xl = admin.get("/api/export/applications?status=신규")
    chk(r_xl.status_code == 200, "STEP 입점신청 엑셀 내보내기 200")
    companies = _xlsx_rows(r_xl, "업체명")
    chk(len(companies) == 2,
        "★★★ 엑셀도 목록과 같은 status=신규 필터를 받아 2건만 담는다"
        "(고치기 전엔 필터 무시하고 4건 전체가 나갔다)", companies)

    # ── [3] 도입문의: 상태='신규' 2건 + 다른 상태 2건, 화면 필터=신규 ────
    for i, st in enumerate(["신규", "신규", "검토중", "승인"]):
        conn.execute(
            "INSERT INTO insurer_leads(company,status,is_demo) "
            f"VALUES('R142문의사{i}',?,0)", (st,))
    conn.commit()
    r_list = admin.get("/api/leads?status=신규")
    n_list = len(r_list.get_json())
    chk(n_list == 2, "★ (사전조건) 도입문의 목록이 status=신규 로 2건만 돌려준다", n_list)
    r_xl = admin.get("/api/export/leads?status=신규")
    chk(r_xl.status_code == 200, "STEP 도입문의 엑셀 내보내기 200")
    orgs = _xlsx_rows(r_xl, "기관명")
    chk(len(orgs) == 2,
        "★★★ 엑셀도 목록과 같은 status=신규 필터를 받아 2건만 담는다"
        "(고치기 전엔 필터 무시하고 4건 전체가 나갔다)", orgs)

    conn.close()
except Exception as e:                                              # noqa: BLE001
    chk(False, "★★★ 라운드142 [1] 엑셀 필터 준비/실행", f"{type(e).__name__} {e}")

# ============================================================
print("\n[3] 실행 — 개인정보 파기: 종료버킷(n4)도 알림본문·조회링크까지 닫힌다")
# ------------------------------------------------------------
try:
    import manage as M

    conn = DB.get_db()
    old_ts = (datetime.now() - timedelta(days=100)).strftime("%Y-%m-%d %H:%M:%S")
    token = "r142token1234567890"
    phone = "01099998888"
    conn.execute(
        "INSERT INTO cases(case_no,status,status_at,created_at,victim_name,victim_phone,"
        "                  track_token,track_issued_at,is_demo) "
        "VALUES('R142-FAIL-1','유찰',?,?,'김철수',?,?,?,0)",
        (old_ts, old_ts, phone, token, old_ts))
    cid = conn.execute("SELECT id FROM cases WHERE case_no='R142-FAIL-1'").fetchone()["id"]
    conn.execute(
        "INSERT INTO notifications(ntype,channel,target,message,created_at,case_id) "
        "VALUES('피해자안내','문자',?,?,?,?)",
        (phone, f"진행조회: https://example.com/track/{token}", old_ts, cid))
    conn.commit()

    rc = M.cmd_prune_personal([])
    chk(rc == 0, "STEP cmd_prune_personal 정상 종료(0)", rc)

    row = conn.execute(
        "SELECT victim_name, victim_phone, track_token FROM cases WHERE id=?", (cid,)).fetchone()
    chk(row["victim_name"] is None and row["victim_phone"] is None,
        "★ (사전조건) 90일 지난 유찰 사건의 피해자 성명·전화가 파기된다", dict(row))
    chk(row["track_token"] is None,
        "★★★ 진행조회 track_token 도 함께 무효화된다"
        "(고치기 전엔 이름·전화만 지워지고 토큰은 그대로 살아 있었다)", dict(row))

    noti = conn.execute(
        "SELECT target, message FROM notifications WHERE case_id=? AND ntype='피해자안내'",
        (cid,)).fetchone()
    chk(noti["target"] == "(파기됨)" and token not in (noti["message"] or ""),
        "★★★ 알림 이력의 전화번호·진행조회 링크 원문도 함께 파기된다"
        "(고치기 전엔 원문 전화번호·링크가 그대로 남아 있었다)", dict(noti))

    import app as A2
    web = A2.app.test_client()
    r = web.get(f"/track/{token}")
    chk(r.status_code != 200,
        "★★★ 파기 뒤에는 그 진행조회 링크가 더는 열리지 않는다(토큰 무효화의 실제 효과)",
        r.status_code)

    conn.close()
except Exception as e:                                              # noqa: BLE001
    chk(False, "★★★ 라운드142 [2] 개인정보 파기 준비/실행", f"{type(e).__name__} {e}")

print("\n" + "=" * 72)
print(f"  통과 {len(PASS)} · 실패 {len(FAIL)}")
if FAIL:
    print("\n  실패 항목:")
    for nm, dt in FAIL:
        print(f"    · {nm}\n        {dt}")
print("=" * 72)
sys.exit(1 if FAIL else 0)
