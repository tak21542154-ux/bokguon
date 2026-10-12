# -*- coding: utf-8 -*-
"""스캔 598라운드 회귀 — 관리자 엑셀/CSV 내보내기 2곳(api_export_xlsx·
   api_settlements_export)에 행수 상한이 전혀 없던 결함.

   조사 배경 — 지시받은 신선한 각도 "대량 데이터 내보내기의 메모리 사용"을
   실제로 재현했다. 형제(api_activity_export·api_partner_onboarding_export)는
   이미 "화면보다 넉넉한 상한 + 잘리면 파일 안에 적는다"를 갖췄는데, cases·
   partners·settlements 등 가장 크게 자라는 6개 표를 한 곳에서 담당하는
   api_export_xlsx 와, 회계 대사용 정산 CSV(api_settlements_export) 두 곳만
   그 방어가 없었다(다른 내보내기 6개는 KPI 집계·1개월/1개 발주처 단위처럼
   자연히 바운드돼 있어 대상이 아니다 — 직접 확인).

   실측(격리 스크래치 DB, settlements 6만행 — 운영 43건의 약 1,400배, cases 도
   조인 대상이라 함께 6만행): 상한 없는 옛 코드는 /api/export/settlements 단일
   요청이 9분 이상 끝나지 않았다(강제 종료 시점 기준 CPU 563초 소진·상주메모리
   512MB 이상 — waitress 스레드 하나를 그만큼 묶어 둔다). 고침 — SQL 문자열은
   그대로 두고(round22가 이미 "conn.execute(sql, p)" 원문을 확인) 두 함수 모두
   fetchmany(cap) 로 자르고 COUNT(*) 서브쿼리로 전체 건수를 따로 구해, 잘렸으면
   그 사실을 **파일 안에** 적는다(활동로그 엑셀과 동일 원칙). cap 기본값·상한은
   형제와 동일(기본 20,000 · 최대 100,000).

     python test_scan_round598_fixes.py
"""
import csv
import io
import os
import shutil
import sys
import tempfile

HERE = os.path.dirname(os.path.abspath(__file__))
PASS, FAIL = [], []


def chk(cond, name, detail=""):
    (PASS if cond else FAIL).append((name, detail))
    print(("  [OK]   " if cond else "  [FAIL] ") + name + ("" if cond else f"   {detail}"))


def read(rel):
    return io.open(os.path.join(HERE, rel), encoding="utf-8").read()


def _py_fn_body(src, fn_name):
    start = src.index(f"\ndef {fn_name}(")
    end = src.find("\ndef ", start + 1)
    return src[start:end if end != -1 else None]


APP_SRC = read("app.py")

print("=" * 72)
print("  스캔 598라운드 회귀 — 내보내기 2곳 행수 상한 없음")
print("=" * 72)

print("\n[1] 소스 확인 — 두 함수 모두 상한·잘림표시를 갖췄다")
_ex = _py_fn_body(APP_SRC, "api_export_xlsx")
_se = _py_fn_body(APP_SRC, "api_settlements_export")

for name, body in (("api_export_xlsx", _ex), ("api_settlements_export", _se)):
    chk('request.args.get("max")' in body, f"★ {name} 이 ?max= 를 받는다(형제와 동일 계약)")
    chk("except (TypeError, ValueError)" in body,
        f"★ {name} 은 max 가 숫자가 아니어도 죽지 않는다")
    chk("min(int(request.args.get(\"max\")" in body and "100000" in body,
        f"★★ {name} 에 상한(무한이 아니다)이 있다 — 최대 100,000")
    chk("fetchmany(cap)" in body, f"★★★★★ {name} 이 fetchall이 아니라 fetchmany(cap)로 자른다"
        "(옛 코드는 상한 없이 전부 메모리에 올렸다)")
    chk("SELECT COUNT(*) AS n FROM (" in body,
        f"★ {name} 이 전체 건수를 별도로 센다(잘렸는지 판정용)")
    chk("을 넘었습니다" in body, f"★★ {name} 이 잘렸을 때 그 사실을 안내 문구로 담는다")

chk("conn.execute(sql, p)" in _ex,
    "★★★ round22 의 원문 검사 문구('conn.execute(sql, p)')가 그대로 살아 있다"
    "(SQL 자체는 안 바꾸고 fetchmany 로만 자른다 — 그 검증을 깨지 않았다)")

print("\n[2] 실행 — 실서버 경로(Flask test_client + 임시 DB)로 실제 재현")
_tmpdir = tempfile.mkdtemp(prefix="round598_")
os.environ["BOKGUON_DB_PATH"] = os.path.join(_tmpdir, "verify.db")
os.environ["BOKGUON_UPLOAD_DIR"] = os.path.join(_tmpdir, "uploads")
os.environ["BOKGUON_ENV"] = "dev"
os.environ["BOKGUON_SEED_DEMO"] = "0"
os.environ["BOKGUON_OTP_POLICY"] = "off"
sys.path.insert(0, HERE)

try:
    import db as DB
    DB.init_db()
    import app as A
    from openpyxl import load_workbook

    N = 15  # 기본 상한(20,000)보다 훨씬 작게 두고 ?max= 로 낮은 상한을 강제해 잘림을 재현한다
    conn = DB.get_db()
    conn.execute("INSERT INTO admins(username,password_hash,name,role) "
                 "VALUES('r598admin','x','검증','최고관리자')")
    for i in range(N):
        conn.execute(
            "INSERT INTO cases(case_no,insurer,category,region,status,created_at,is_demo) "
            "VALUES(?,?,?,?,?,datetime('now'),0)",
            (f"R598-{i:03d}", "R598보험", "누수", "서울", "완료"))
        conn.execute(
            "INSERT INTO settlements(case_no,insurer,partner_name,amount,fee,payout,"
            "                        created_at,is_demo) VALUES(?,?,?,?,?,?,datetime('now'),0)",
            (f"R598-{i:03d}", "R598보험", f"R598업체{i}", 1000000, 100000, 900000))
    conn.commit()
    conn.close()

    admin = A.app.test_client()
    with admin.session_transaction() as sess:
        sess["admin"] = "r598admin"; sess["admin_role"] = "최고관리자"; sess["admin_name"] = "검증"

    print("\n  [2-1] xlsx(api_export_xlsx) — 상한보다 낮은 ?max= 로 잘림 강제")
    r = admin.get("/api/export/cases?max=6")
    chk(r.status_code == 200, "STEP cases 엑셀(?max=6) 200", r.status_code)
    if r.status_code == 200:
        wb = load_workbook(io.BytesIO(r.data))
        ws = wb.active
        row1 = [c.value for c in ws[1]]
        row2 = [c.value for c in ws[2]]
        chk(str(row1[0] or "").startswith("조건에 맞는"),
            "★★★★★ 잘렸으면 1행에 안내 문구가 들어간다(파일 안에 남는다)", row1)
        chk(f"{N}" in str(row1[0]) and "6" in str(row1[0]),
            "★★ 안내 문구에 전체건수(15)와 담긴건수(6)가 함께 적힌다", row1)
        chk(row2[0] == "접수번호", "★ 2행이 실제 머리글이다(잘렸을 때만 1행이 밀린다)", row2)
        n_data = ws.max_row - 2
        chk(n_data == 6, "★★★★★ 실제로 요청한 상한(6)만큼만 담겼다(15건 중)", n_data)

    print("\n  [2-2] xlsx — 상한 밖(15<20000, 기본 상한)이면 잘리지 않고 헤더가 1행 그대로")
    r2 = admin.get("/api/export/cases")
    chk(r2.status_code == 200, "STEP cases 엑셀(기본) 200", r2.status_code)
    if r2.status_code == 200:
        wb2 = load_workbook(io.BytesIO(r2.data))
        ws2 = wb2.active
        chk([c.value for c in ws2[1]][0] == "접수번호",
            "★★★ 안 잘리면 기존 계약(헤더=1행)이 그대로다(회귀 없음)",
            [c.value for c in ws2[1]])
        chk(ws2.max_row - 1 == N, "★ 15건이 그대로 다 나온다", ws2.max_row - 1)

    print("\n  [2-3] xlsx — max 가 숫자가 아니어도 500 이 아니다")
    r3 = admin.get("/api/export/cases?max=abc")
    chk(r3.status_code == 200, "STEP ?max=abc 도 200(기본 상한으로 대체)", r3.status_code)

    print("\n  [2-4] csv(api_settlements_export) — 상한보다 낮은 ?max= 로 잘림 강제")
    r4 = admin.get("/api/settlements/export?max=4")
    chk(r4.status_code == 200, "STEP 정산 CSV(?max=4) 200", r4.status_code)
    if r4.status_code == 200:
        body_txt = r4.get_data(as_text=True).lstrip("﻿")
        csv_rows = list(csv.reader(io.StringIO(body_txt)))
        chk(csv_rows[0][0].startswith("※ 조건에 맞는"),
            "★★★★★ CSV 도 잘렸으면 1행에 안내 문구가 들어간다", csv_rows[0])
        chk(csv_rows[1][0] == "구분", "★ 2행이 실제 머리글이다", csv_rows[1])
        n_data_csv = len(csv_rows) - 2
        chk(n_data_csv == 4, "★★★★★ 정산 CSV 도 요청한 상한(4)만큼만 담겼다(15건 중)", n_data_csv)

    print("\n  [2-5] csv — 안 잘리면 옛 계약대로 안내 문구 없이 머리글이 1행")
    r5 = admin.get("/api/settlements/export")
    chk(r5.status_code == 200, "STEP 정산 CSV(기본) 200", r5.status_code)
    if r5.status_code == 200:
        body5 = r5.get_data(as_text=True).lstrip("﻿")
        csv_rows5 = list(csv.reader(io.StringIO(body5)))
        chk(csv_rows5[0][0] == "구분",
            "★★★ 안 잘리면 회귀 없이 1행이 바로 머리글이다", csv_rows5[0])
        chk(len(csv_rows5) - 1 == N, "★ 15건이 그대로 다 나온다", len(csv_rows5) - 1)

    print("\n  [2-6] 내보낸 사실이 '잘렸다'는 표시와 함께 감사로그에 남는다")
    #  ⚠ 각 action 의 **첫** 기록이 잘린 호출(2-1·2-4)이다 — 뒤이은 2-2/2-5(안 잘림)가
    #    로그를 더 남기므로 반드시 id ASC 로 첫 건을 봐야 한다(DESC는 나중 호출을 본다).
    conn = DB.get_db()
    first_xlsx_log = conn.execute(
        "SELECT detail FROM activity_log WHERE action='엑셀 내보내기' "
        "ORDER BY id ASC LIMIT 1").fetchone()
    first_csv_log = conn.execute(
        "SELECT detail FROM activity_log WHERE action='정산 내역 내보내기' "
        "ORDER BY id ASC LIMIT 1").fetchone()
    conn.close()
    chk(bool(first_xlsx_log) and "상한 초과로 잘림" in first_xlsx_log[0],
        "★★ 잘린 xlsx 내보내기는 감사로그에도 '상한 초과로 잘림'이 남는다",
        first_xlsx_log[0] if first_xlsx_log else None)
    chk(bool(first_csv_log) and "상한 초과로 잘림" in first_csv_log[0],
        "★★ 잘린 정산 CSV 내보내기도 감사로그에 '상한 초과로 잘림'이 남는다",
        first_csv_log[0] if first_csv_log else None)

    print("\n[3] 되돌림-증명 — 상한 로직이 없으면 이 테스트가 실제로 잡아내는가")
    #  ⚠ 디스크의 app.py 는 절대 건드리지 않는다 — Flask 의 view_functions
    #    딕셔너리 항목만 임시로 바꿔치기하고 finally 에서 반드시 원복한다.
    _orig_view = A.app.view_functions["api_export_xlsx"]

    def _old_broken_view(section):
        """598라운드 이전 동작 재현 — 상한 없이 그대로 fetchall."""
        if section not in A.EXPORT_SECTIONS:
            return A.jsonify({"ok": False, "error": "내보낼 수 없는 항목입니다."}), 404
        title, sql = A.EXPORT_SECTIONS[section]
        sql2 = sql.replace("{WHERE}", "WHERE COALESCE(is_demo,0)=0") if "{WHERE}" in sql else sql
        conn2 = A.DB.get_db()
        cur2 = conn2.execute(sql2, [])
        headers2 = [d[0] for d in cur2.description]
        data2 = cur2.fetchall()          # ← 598라운드 이전: 상한 없음
        conn2.close()
        from openpyxl import Workbook as _WB
        wb2 = _WB(); ws3 = wb2.active
        ws3.append(headers2)
        for rr in data2:
            ws3.append([rr[h] for h in headers2])
        buf2 = io.BytesIO(); wb2.save(buf2); buf2.seek(0)
        return A.make_response(buf2.read())

    A.app.view_functions["api_export_xlsx"] = _old_broken_view
    try:
        r_old = admin.get("/api/export/cases?max=6")   # 옛 코드는 이 max=6 을 아예 모른다
        wb_old = load_workbook(io.BytesIO(r_old.data))
        n_old = wb_old.active.max_row - 1
        chk(n_old == N,
            "★★★★★ 상한 로직을 빼면(구버전 재현) ?max=6 을 무시하고 15건을 전부 돌려준다"
            " — 이 테스트가 실제로 회귀를 잡아낸다는 증거", n_old)
    finally:
        A.app.view_functions["api_export_xlsx"] = _orig_view

    #  원복 후 같은 요청이 다시 정상(잘림)으로 돌아오는지 — 바꿔치기가 새지 않았다
    r_restored = admin.get("/api/export/cases?max=6")
    wb_restored = load_workbook(io.BytesIO(r_restored.data))
    chk(wb_restored.active.max_row - 2 == 6,
        "★ 원복 후에는 다시 상한(6)이 정상 적용된다(바꿔치기가 안 샜다)",
        wb_restored.active.max_row - 2)

except Exception as e:                                              # noqa: BLE001
    chk(False, "★★★★★ 라운드598 서버측 준비/실행", f"{type(e).__name__} {e}")
finally:
    os.environ.pop("BOKGUON_DB_PATH", None)
    os.environ.pop("BOKGUON_UPLOAD_DIR", None)
    os.environ.pop("BOKGUON_ENV", None)
    os.environ.pop("BOKGUON_SEED_DEMO", None)
    os.environ.pop("BOKGUON_OTP_POLICY", None)
    shutil.rmtree(_tmpdir, ignore_errors=True)

print("\n" + "=" * 72)
print(f"  통과 {len(PASS)} · 실패 {len(FAIL)}")
if FAIL:
    print("\n  실패 항목:")
    for nm, dt in FAIL:
        print(f"    - {nm}\n        {dt}")
print("=" * 72)
sys.exit(1 if FAIL else 0)
